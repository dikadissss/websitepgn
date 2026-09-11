"""PDE workbook and Peta Harian page of a daily report."""
import datetime
from copy import copy
from io import BytesIO
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins
from PIL import Image, ImageDraw

from core.dates import MONTHS

from .maps import DEPTH_CLASSES, DEPTH_LABELS, MAGNITUDE_LABELS, font, render_map

STATIC_DIRECTORY = Path(__file__).resolve().parent / 'static' / 'daily_report'
PDE_TEMPLATE = STATIC_DIRECTORY / 'PDE.xlsx'
LOGO_PATH = STATIC_DIRECTORY / 'bmkg_logo.png'
ENGLISH_MONTHS = ('January', 'February', 'March', 'April', 'May', 'June', 'July',
                  'August', 'September', 'October', 'November', 'December')

PDE_FIRST_ROW = 13  # Row 13 of the template holds the style of the event rows.
PDE_COLUMNS = 9
SIGNATURE_FONT = Font(name='Arial Narrow', size=11)
SIGNATURE_NAME_FONT = Font(name='Arial Narrow', size=11, bold=True, underline='single')
SIGNATURE_LINE = Border(bottom=Side(style='thin'))
CENTERED = Alignment(horizontal='center', vertical='center')

PAGE_SIZE = (1754, 1240)  # A4 landscape at 150 dpi.
PAGE_DPI = 150
BOX_FILL = (253, 200, 150)
BOX_LEFT, BOX_RIGHT = 337, 1417
LEGEND_CIRCLE_RADII = (7, 10, 13)
# The map page in the Excel export: A4 landscape inside 0.4 in margins, in pixels at 96 dpi (openpyxl's unit).
MAP_SHEET_IMAGE_SIZE = (1005, 710)
MAP_SHEET_PRINT_AREA = 'A1:P36'  # The cells under the image (default column width 64 px, row height 20 px).


def english_date(date):
    """'September 07, 2026'"""
    return f'{ENGLISH_MONTHS[date.month - 1]} {date:%d}, {date.year}'


def signing_date(date):
    """'08 September 2026'"""
    return f'{date:%d} {MONTHS[date.month - 1]} {date.year}'


def summary_sentence(report):
    table = report.event_table()
    day = english_date(report.report_date)
    magnitudes = pd.to_numeric(table['Mag'], errors='coerce')
    if magnitudes.isna().all():
        return f'On this day of {day} no earthquake was recorded.'
    largest = table.loc[magnitudes.idxmax()]
    return (f'On this day of {day} has occured {len(table)} earthquakes with the largest magnitude was '
            f'{magnitudes.max():.1f} {largest["TypeMag"]} in {largest["Region"]}')


def export_documents(report):
    """Pages of the daily report in print order: the map, then the PDE."""
    return [build_map_pdf(report), build_pde_workbook(report)]


def report_filename(report):
    return f'Daily_Report_{report.report_date:%Y-%m-%d}'


def build_report_workbook(report):
    """Excel export of the report: the Peta Harian page as the first sheet, then the PDE."""
    workbook = build_pde_workbook(report)
    sheet = workbook.create_sheet('Peta', 0)
    buffer = BytesIO()
    _map_page(report).save(buffer, 'PNG')
    image = SheetImage(buffer)
    image.width, image.height = MAP_SHEET_IMAGE_SIZE
    sheet.add_image(image, 'A1')
    sheet.sheet_view.showGridLines = False
    sheet.print_area = MAP_SHEET_PRINT_AREA
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.4)
    workbook.active = 0
    return workbook


# PDE workbook

def build_pde_workbook(report):
    workbook = openpyxl.load_workbook(PDE_TEMPLATE)
    sheet = workbook.active
    sheet.title = f'{report.report_date:%d-%m-%Y}'
    sheet['A10'] = f'Preliminary Determination of Epicenter, {english_date(report.report_date)}'

    styles = [copy(sheet.cell(row=PDE_FIRST_ROW, column=column)._style) for column in range(1, PDE_COLUMNS + 1)]
    events = report.event_table().to_dict(orient='records')
    for number, event in enumerate(events, 1):
        magnitude = event['Mag']
        values = (
            number,
            datetime.datetime.strptime(event['Date'], '%Y-%m-%d'),
            datetime.time.fromisoformat(event['OT (UTC)']),
            event['Lat'], event['Long'], event['D(Km)'],
            None if pd.isna(magnitude) else float(magnitude),
            event['TypeMag'], event['Region'],
        )
        sheet_row = PDE_FIRST_ROW + number - 1
        for column, (value, style) in enumerate(zip(values, styles), 1):
            cell = sheet.cell(row=sheet_row, column=column, value=value)
            cell._style = copy(style)
        sheet.row_dimensions[sheet_row].height = sheet.row_dimensions[PDE_FIRST_ROW].height

    last_row = PDE_FIRST_ROW + max(len(events), 1) - 1
    signature_row = _add_signatures(sheet, report, last_row + 2)
    sheet.print_area = f'A1:I{signature_row}'
    return workbook


def _add_signatures(sheet, report, row):
    """'Jakarta, <date>', 'Petugas onduty' / 'Mengetahui', then the officer's name and NIP; returns the last row used.

    "Mengetahui" is signed by hand: its name and NIP stay empty, above a signature line as on the map page."""
    blocks = (
        (row, 'G', 'I', f'Jakarta, {signing_date(report.date)}', SIGNATURE_FONT),
        (row + 1, 'B', 'E', 'Petugas onduty', SIGNATURE_FONT),
        (row + 1, 'G', 'I', 'Mengetahui', SIGNATURE_FONT),
        (row + 5, 'B', 'E', report.operator.name, SIGNATURE_NAME_FONT),
        (row + 6, 'B', 'E', f'NIP. {report.operator.NIP}', SIGNATURE_FONT),
    )
    for sheet_row, first, last, text, cell_font in blocks:
        cell = sheet[f'{first}{sheet_row}']
        cell.value, cell.font, cell.alignment = text, cell_font, CENTERED
        sheet.merge_cells(f'{first}{sheet_row}:{last}{sheet_row}')
    for column in 'GHI':
        sheet[f'{column}{row + 5}'].border = SIGNATURE_LINE
    return row + 6


# Peta Harian page

def build_map_pdf(report):
    buffer = BytesIO()
    _map_page(report).save(buffer, 'PDF', resolution=PAGE_DPI)
    return buffer.getvalue()


def _wrap(draw, text, text_font, width):
    lines, line = [], ''
    for word in text.split():
        candidate = f'{line} {word}'.strip()
        if line and draw.textlength(candidate, font=text_font) > width:
            lines.append(line)
            line = word
        else:
            line = candidate
    return lines + [line] if line else lines


def _box(draw, box):
    draw.rounded_rectangle(box, radius=8, fill=BOX_FILL, outline='black', width=2)


def _map_page(report):
    if not report.map_path.exists():
        render_map(report.event_table(), report.map_path)
    page = Image.new('RGB', PAGE_SIZE, 'white')
    draw = ImageDraw.Draw(page)
    draw.rectangle((BOX_LEFT - 10, 40, BOX_RIGHT + 10, 1006), outline='black', width=2)

    _box(draw, (BOX_LEFT, 50, BOX_RIGHT, 130))
    center = (BOX_LEFT + BOX_RIGHT) / 2
    draw.text((center, 72), 'MAP OF EARTHQUAKE EPICENTER', fill='black', font=font(30, bold=True), anchor='mm')
    draw.text((center, 108), english_date(report.report_date).upper(), fill='black', font=font(30, bold=True), anchor='mm')

    with Image.open(report.map_path) as map_image:
        map_width = BOX_RIGHT - BOX_LEFT
        map_height = round(map_width * map_image.height / map_image.width)
        page.paste(map_image.convert('RGB').resize((map_width, map_height), Image.LANCZOS), (BOX_LEFT, 140))

    boxes_top, boxes_bottom = 846, 996
    _legend(draw, (BOX_LEFT, boxes_top, 767, boxes_bottom))
    _source(page, draw, report, (777, boxes_top, 1127, boxes_bottom))
    _summary(draw, report, (1137, boxes_top, BOX_RIGHT, boxes_bottom))
    _page_signatures(draw, report)
    return page


def _legend(draw, box):
    _box(draw, box)
    left, top, right, bottom = box
    label_width = 105
    column_width = (right - left - label_width) / 3
    header_height = 34
    row_height = (bottom - top - header_height) / 3
    small = font(13)
    for index in range(1, 4):
        x = left + label_width + (index - 1) * column_width
        draw.line((x, top, x, top + header_height), fill='black', width=1)
    draw.line((left, top + header_height, right, top + header_height), fill='black', width=1)
    draw.text((left + label_width / 2, top + header_height / 2), 'Mag', fill='black', font=small, anchor='mm')
    for index, label in enumerate(DEPTH_LABELS):
        x = left + label_width + (index + 0.5) * column_width
        draw.text((x, top + header_height / 2), label, fill='black', font=small, anchor='mm')
    for row, (label, radius) in enumerate(zip(MAGNITUDE_LABELS, LEGEND_CIRCLE_RADII)):
        y = top + header_height + (row + 0.5) * row_height
        draw.text((left + label_width / 2, y), label, fill='black', font=small, anchor='mm')
        for index, (_, color) in enumerate(DEPTH_CLASSES):
            x = left + label_width + (index + 0.5) * column_width
            draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill=color, outline='black', width=1)


def _source(page, draw, report, box):
    _box(draw, box)
    left, top, right, bottom = box
    with Image.open(LOGO_PATH) as logo:
        logo = logo.convert('RGBA')
        logo.thumbnail((110, bottom - top - 24), Image.LANCZOS)
        page.paste(logo, (left + 14, round((top + bottom - logo.height) / 2)), logo)
        text_left = left + 14 + logo.width + 16
    text = (f'Source: Earthquake Data {english_date(report.report_date)}, '
            'Earthquake and Tsunami Information Division BMKG')
    _paragraph(draw, text, (text_left, top, right - 12, bottom), font(16))


def _summary(draw, report, box):
    _box(draw, box)
    left, top, right, bottom = box
    _paragraph(draw, summary_sentence(report), (left + 12, top, right - 12, bottom), font(17))


def _paragraph(draw, text, box, text_font):
    """Text wrapped to the box width and centered vertically, left aligned."""
    left, top, right, bottom = box
    lines = _wrap(draw, text, text_font, right - left)
    line_height = text_font.size * 1.3
    y = (top + bottom) / 2 - line_height * len(lines) / 2
    for line in lines:
        draw.text((left, y), line, fill='black', font=text_font)
        y += line_height


def _page_signatures(draw, report):
    left_center, right_center = 480, 1270
    text_font, name_font = font(19), font(19, bold=True)
    draw.text((right_center, 1025), f'Jakarta, {signing_date(report.date)}', fill='black', font=text_font, anchor='mm')
    draw.text((left_center, 1052), 'Petugas Onduty', fill='black', font=text_font, anchor='mm')
    draw.text((right_center, 1052), 'Mengetahui', fill='black', font=text_font, anchor='mm')
    for center in (left_center, right_center):
        draw.line((center - 160, 1160, center + 160, 1160), fill='black', width=2)
    # "Mengetahui" is signed by hand: only the officer's name and NIP are printed.
    draw.text((left_center, 1170), report.operator.name, fill='black', font=name_font, anchor='ma')
    draw.text((left_center, 1198), f'NIP. {report.operator.NIP}', fill='black', font=text_font, anchor='ma')
