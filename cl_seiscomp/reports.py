import datetime
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image

from core.dates import date_from_code, date_range_to_string, format_date_indonesian, get_hari_indonesia

TEMPLATE_PATH = Path(__file__).resolve().parent / 'static' / 'cl_seiscomp' / 'cl_seiscomp.xlsx'
SLMON_IMAGE_SIZE = (8.6 * 96, 4.14 * 96)  # Width and height in pixels (inches * 96 dpi).


def export_documents(record):
    """Documents printed for the record, in order (see core.duty.JOBS)."""
    return [build_workbook(record)]


def build_workbook(record):
    workbook = openpyxl.load_workbook(TEMPLATE_PATH)
    record_date = date_from_code(record.cs_id)
    # The Malam checklist runs past midnight, so its report covers two days.
    is_night = record.shift.upper() == 'MALAM'
    _fill_checklist(workbook['checklist_seiscomp'], record, record_date, is_night)
    _fill_slmon(workbook['slmon'], record, record_date + datetime.timedelta(days=1) if is_night else record_date)
    return workbook


def _fill_checklist(sheet, record, record_date, is_night):
    sheet.title = 'Checklist Seiscomp'
    if is_night:
        sheet['R3'] = date_range_to_string(record_date, record_date + datetime.timedelta(days=1))
    else:
        sheet['R3'] = f'{get_hari_indonesia(record_date)}, {format_date_indonesian(record_date)}'
    sheet['A3'] = f'KELOMPOK: {record.kelompok}'
    sheet['A2'] = f'SHIFT {record.shift.upper()}'
    sheet['H286'] = f'{record.operator}'
    sheet['D5'] = sheet['P5'] = sheet['H276'] = f'JAM {record.jam_pelaksanaan}'

    stations = [set((text or '').splitlines()) for text in (record.gaps, record.spikes, record.blanks)]
    # Station codes are listed in two blocks; mark gaps/spikes/blanks next to each code.
    _mark_stations(sheet, code_column=2, last_row=287, mark_columns=(4, 5, 6), stations=stations)
    _mark_stations(sheet, code_column=9, last_row=274, mark_columns=(16, 17, 18), stations=stations)


def _mark_stations(sheet, code_column, last_row, mark_columns, stations):
    for row in range(7, last_row + 1):
        code = sheet.cell(row=row, column=code_column).value
        for column, codes in zip(mark_columns, stations):
            if code in codes:
                sheet.cell(row=row, column=column).value = 1


def _fill_slmon(sheet, record, report_date):
    tanggal = format_date_indonesian(report_date)
    sheet['A2'] = f'{tanggal}, pukul {record.jam_pelaksanaan}'
    sheet['M24'] = f'Jakarta, {tanggal}'
    sheet['C28'] = f'{record.operator}'

    if not record.slmon_image:
        sheet['B3'] = 'No image'
        return
    try:
        image = Image(record.slmon_image.path)
    except FileNotFoundError:
        sheet['B3'] = 'Image file not found'
        return
    image.anchor = 'B3'
    image.width, image.height = SLMON_IMAGE_SIZE
    if sheet._images:
        sheet._images[0] = image
    else:
        sheet.add_image(image)
