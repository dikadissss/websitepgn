"""Printed Checklist Tide Gauge, drawn with openpyxl in the layout of the InaTNT recap sheet
(Rekap_InaTNT_yyyymmdd_hhmm.xlsx of the old tools/Cheklist_TG_2026.py): a table per network, split into a left
(A-G) and a right (I-O) block, then the counts per network and the signatures."""
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from bast.reports import group_label
from core.dates import format_date_indonesian
from core.reports import THIN_BORDER

from . import feed

BLOCK_WIDTHS = (5.71, 15.71, 35.71, 18.71, 8.71, 8.71, 8.71)  # No., Stasiun, Lokasi, Provinsi, Gaps, Spike, Blank.
LEFT_BLOCK, RIGHT_BLOCK = 1, 9  # First columns (A, I) of the two blocks; column H between them is narrow.
LAST_COLUMN = 15
FIRST_TABLE_ROW = 8
SPLIT_AT = {'WL': 17, 'IC': 16}  # Rows of the left block; other networks are halved when longer than HALVE_ABOVE.
HALVE_ABOVE = 45
MARKS = ('gaps', 'spike', 'blank')  # The Monitor columns.
CHECK = '✓'
CHECKED, UNCHECKED = '☑', '☐'
GREY = PatternFill(fill_type='solid', start_color='FFD9D9D9', end_color='FFD9D9D9')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center')


def font(size=11, bold=False, **options):
    return Font(name='Calibri', size=size, bold=bold, **options)


HEADER = {'style': font(bold=True), 'alignment': CENTER, 'border': True, 'fill': GREY}


def export_documents(record):
    """Documents printed for the record, in order (see core.duty.JOBS)."""
    return [build_workbook(record)]


def build_workbook(record):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = record.code
    for offset, width in enumerate(BLOCK_WIDTHS):
        for first in (LEFT_BLOCK, RIGHT_BLOCK):
            sheet.column_dimensions[get_column_letter(first + offset)].width = width
    sheet.column_dimensions['H'].width = 2.71

    _header(sheet, record)
    row = FIRST_TABLE_ROW
    for blocks in _sections(feed.group_stations(record.stations)):
        row = _table(sheet, row, blocks, record.web_accessible)
    row = _summary(sheet, row + 1, record.summary())
    last_row = _signatures(sheet, row + 3, record)

    sheet.print_area = f'A1:{get_column_letter(LAST_COLUMN)}{last_row}'
    sheet.page_setup.orientation = 'portrait'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0  # As many pages as needed.
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.3, right=0.3, top=0.5, bottom=0.5)
    return workbook


def _write(sheet, row, column, value, *, to_column=None, to_row=None, style=None, alignment=LEFT, border=False,
           fill=None):
    """Value in the cell, merged with the cells up to (to_row, to_column) when given."""
    to_column, to_row = to_column or column, to_row or row
    for sheet_row in range(row, to_row + 1):
        for sheet_column in range(column, to_column + 1):
            cell = sheet.cell(sheet_row, sheet_column)
            if border:
                cell.border = THIN_BORDER
            if fill:
                cell.fill = fill
    cell = sheet.cell(row, column)
    cell.value, cell.font, cell.alignment = value, style or font(), alignment
    if (row, column) != (to_row, to_column):
        sheet.merge_cells(start_row=row, start_column=column, end_row=to_row, end_column=to_column)


def _header(sheet, record):
    _write(sheet, 1, 1, record.title, to_column=LAST_COLUMN, style=font(14, True), alignment=CENTER)
    bold = font(bold=True)
    for row, column, label, value in (
            (3, LEFT_BLOCK, 'Kelompok :', group_label(record.kelompok)),
            (4, LEFT_BLOCK, 'Jadwal Shift :', record.shift),
            (3, RIGHT_BLOCK, 'Hari / Tanggal :', format_date_indonesian(record.date).upper()),
            (4, RIGHT_BLOCK, 'Waktu Pengecekan :', f'{record.check_time:%H:%M} WIB')):
        _write(sheet, row, column, label, style=bold)
        _write(sheet, row, column + 2, value, style=bold)


def _sections(groups):
    """The tables in print order, each a list of blocks (first column, network, stations, first number, heading)."""
    for net, stations in groups.items():
        if net == 'BY' and 'ID' in groups:
            continue
        if net == 'ID':  # IDSL and InaBuoy are printed next to each other.
            blocks = [(LEFT_BLOCK, net, stations, 1, True)]
            if 'BY' in groups:
                blocks.append((RIGHT_BLOCK, 'BY', groups['BY'], 1, True))
            yield blocks
            continue
        split = SPLIT_AT.get(net, -(-len(stations) // 2) if len(stations) > HALVE_ABOVE else len(stations))
        blocks = [(LEFT_BLOCK, net, stations[:split], 1, True)]
        if stations[split:]:
            blocks.append((RIGHT_BLOCK, net, stations[split:], split + 1, False))
        yield blocks


def _access_line(accessible):
    yes, no = (CHECKED, UNCHECKED) if accessible else (UNCHECKED, CHECKED)
    return f'Keterangan: Web Tide Gauge bisa di Akses :   {yes} Ya     {no} Tidak'


def _table(sheet, row, blocks, accessible):
    """Heading, 'Web bisa di Akses' line, two header rows and a row per station; returns the next table's row."""
    for column, net, stations, first_number, heading in blocks:
        if heading:
            _write(sheet, row, column, feed.heading(net), style=font(bold=True, italic=True, underline='single'))
            # DejaVu Sans has the boxes: with Calibri, LibreOffice falls back to a color emoji font.
            _write(sheet, row + 1, column, _access_line(accessible), style=Font(name='DejaVu Sans', size=10))
        _table_header(sheet, row + 3, column, net)
        for offset, station in enumerate(stations):
            _station(sheet, row + 5 + offset, column, first_number + offset, net, station)
    return row + max(len(stations) for _, _, stations, _, _ in blocks) + 6


def _table_header(sheet, row, column, net):
    _write(sheet, row, column, 'No.', to_row=row + 1, **HEADER)
    _write(sheet, row, column + 1, 'STASIUN', to_row=row + 1, **HEADER)
    if feed.has_region(net):
        _write(sheet, row, column + 2, 'LOKASI', to_row=row + 1, **HEADER)
        _write(sheet, row, column + 3, feed.region_header(net), to_row=row + 1, **HEADER)
    else:
        _write(sheet, row, column + 2, 'LOKASI', to_column=column + 3, to_row=row + 1, **HEADER)
    _write(sheet, row, column + 4, 'Monitor', to_column=column + 6, **HEADER)
    for offset, label in enumerate(('Gaps', 'Spike', 'Blank')):
        _write(sheet, row + 1, column + 4 + offset, label, **HEADER)


def _station(sheet, row, column, number, net, station):
    _write(sheet, row, column, number, alignment=CENTER, border=True)
    _write(sheet, row, column + 1, station['code'], border=True)
    if feed.has_region(net):
        _write(sheet, row, column + 2, station['location'], border=True)
        _write(sheet, row, column + 3, station['region'], border=True)
    else:
        _write(sheet, row, column + 2, station['location'], to_column=column + 3, border=True)
    for offset, status in enumerate(MARKS):
        _write(sheet, row, column + 4 + offset, CHECK if station['status'] == status else None, alignment=CENTER,
               border=True)


# Columns of the summary: No., Sistem Platform, Total Alat, Gaps, Spike, Blank, Normal.
SUMMARY_COLUMNS = ((1, 1), (2, 4), (5, 6), (7, 7), (8, 9), (10, 10), (11, 11))
COUNTS = ('total', 'gaps', 'spike', 'blank', 'normal')


def _summary_row(sheet, row, values, columns=SUMMARY_COLUMNS, **style):
    for (first, last), value in zip(columns, values):
        _write(sheet, row, first, value, to_column=last, **style)


def _summary(sheet, row, summary):
    """Counts per network, TOTAL KESELURUHAN and PERSENTASE; returns the last row."""
    _summary_row(sheet, row, ('No.', 'Sistem Platform', 'Total Alat', 'Gaps', 'Spike', 'Blank', 'Normal'), **HEADER)
    for number, network in enumerate(summary['networks'], 1):
        row += 1
        _write(sheet, row, 1, number, alignment=CENTER, border=True)
        _write(sheet, row, 2, network['heading'], to_column=4, border=True)
        _summary_row(sheet, row, [network[count] for count in COUNTS], SUMMARY_COLUMNS[2:], alignment=CENTER,
                     border=True)
    total_columns = ((1, 4), *SUMMARY_COLUMNS[2:])
    _summary_row(sheet, row + 1, ['TOTAL KESELURUHAN', *(summary['total'][count] for count in COUNTS)],
                 total_columns, **HEADER)
    _summary_row(sheet, row + 2, ['PERSENTASE', '100%', *(summary['percent'][count] for count in COUNTS[1:])],
                 total_columns, **HEADER)
    return row + 2


def _signatures(sheet, row, record):
    """'Operator on Duty' with the operator's name, and 'Mengetahui, Pejabat on Duty' signed by hand."""
    bold = font(bold=True)
    _write(sheet, row, 11, f'Jakarta, {format_date_indonesian(record.date)}', to_column=14, alignment=CENTER)
    _write(sheet, row + 1, 11, 'Mengetahui,', to_column=14, alignment=CENTER)
    _write(sheet, row + 2, 2, 'Operator on Duty', to_column=5, style=bold, alignment=CENTER)
    _write(sheet, row + 2, 11, 'Pejabat on Duty', to_column=14, style=bold, alignment=CENTER)
    _write(sheet, row + 6, 2, f'( {record.operator.name} )', to_column=5, alignment=CENTER)
    _write(sheet, row + 6, 11, '(.......................................)', to_column=14, alignment=CENTER)
    return row + 6
