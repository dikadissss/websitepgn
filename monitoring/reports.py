"""Printed form of a monitoring checklist, built with openpyxl in the layout of the 2026 templates
(Checklist Email_2026.xlsx and Checklist Sirine_2026.xlsx): header, one table per catalog section, signatures."""
import math

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.page import PageMargins

from bast.reports import group_label
from core.dates import format_date_indonesian, get_hari_indonesia
from core.reports import THIN_BORDER

from .answers import record_sections
from .models import ChecklistSection

EMAIL_WEB_WIDTHS = (4.55, 12.66, 18.44, 58.33, 6.66, 56.66, 7.11, 27.66)  # Columns A-H of the templates.
DEVICE_WIDTHS = (7.45, 60.33, 21.66, 31.55)  # Columns A-D.
CHECK = '✓'  # Ya / Tidak of the device checklist.
CHECKED, UNCHECKED = '☑', '☐'  # Boxes of the Y/N and J/BJ (B/BB) options of the email/web tables.
HEADER_FILL = PatternFill(fill_type='solid', start_color='FF948A54', end_color='FF948A54')
WHITE = 'FFFFFFFF'
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
LEFT_TOP = Alignment(horizontal='left', vertical='top', wrap_text=True)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
SHRINK = Alignment(horizontal='left', vertical='center', shrink_to_fit=True)


def font(size, bold=False, **options):
    return Font(name='Calibri', size=size, bold=bold, **options)


def export_documents(record):
    """Documents printed for the record, in order (see core.duty.JOBS)."""
    return [build_workbook(record)]


def build_workbook(record):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = record.code
    sheet.sheet_view.showGridLines = False
    if record.form == ChecklistSection.Form.EMAIL_WEB:
        _email_web_sheet(sheet, record)
    else:
        _device_sheet(sheet, record)
    return workbook


def _span(first, last):
    return [get_column_letter(index)
            for index in range(column_index_from_string(first), column_index_from_string(last) + 1)]


def _write(sheet, row, first, last, value, style=None, alignment=LEFT, border=False, fill=None, last_row=None):
    """Value in the cells first..last of the row (down to last_row) in the font `style`, merged when they are more
    than one."""
    last_row = last_row or row
    for sheet_row in range(row, last_row + 1):
        for column in _span(first, last):
            cell = sheet[f'{column}{sheet_row}']
            if border:
                cell.border = THIN_BORDER
            if fill:
                cell.fill = fill
    cell = sheet[f'{first}{row}']
    cell.value, cell.alignment, cell.font = value, alignment, style or font(11)
    if (first, row) != (last, last_row):
        sheet.merge_cells(f'{first}{row}:{last}{last_row}')


def _outline(sheet, first_row, last_row, first, last, style='thin'):
    """A border around the block, without lines inside it."""
    side = Side(style=style)
    columns = _span(first, last)
    for row in range(first_row, last_row + 1):
        for column in columns:
            sheet[f'{column}{row}'].border = Border(
                left=side if column == first else None, right=side if column == last else None,
                top=side if row == first_row else None, bottom=side if row == last_row else None)


def _width(sheet, first, last):
    return sum(sheet.column_dimensions[column].width for column in _span(first, last))


def _lines(text, width, size):
    """Lines of the text wrapped in a cell `width` wide (the width unit is about one Calibri 11 character)."""
    characters = max(width - 1, 1) * 11 / size
    return sum(max(1, math.ceil(len(part) / characters)) for part in str(text or '').split('\n'))


def _fit_height(sheet, rows, texts, minimum):
    """Heights of the rows sharing the wrapped texts [(text, width, font size)]: the PDF keeps the file's heights."""
    needed = max((_lines(text, width, size) * size * 1.3 + 4 for text, width, size in texts), default=0)
    for row in rows:
        sheet.row_dimensions[row].height = max(minimum, needed / len(rows))


def _columns(sheet, widths):
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _day(date):
    return f'{get_hari_indonesia(date)}, {format_date_indonesian(date)}'


def _page_setup(sheet, last_column, last_row, margins):
    """A4 portrait on one page, as the templates are printed."""
    sheet.print_area = f'A1:{last_column}{last_row}'
    sheet.page_setup.orientation = 'portrait'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.page_margins = margins


def _notes(sheet, section, row, left, right):
    """The section's notes under its table: `note` in the columns left, `note_right` in the columns right."""
    if not (section['note'] or section['note_right']):
        return row
    _write(sheet, row, *left, section['note'] or None, font(11), LEFT_TOP)
    _write(sheet, row, *right, section['note_right'] or None, font(11), LEFT_TOP)
    _fit_height(sheet, [row], [(section['note'], _width(sheet, *left), 11),
                               (section['note_right'], _width(sheet, *right), 11)], 15)
    return row + 1


# Monitoring Email, Web, Medsos dan WRS-NG

def _email_web_sheet(sheet, record):
    _columns(sheet, EMAIL_WEB_WIDTHS)
    _write(sheet, 1, 'A', 'H', record.title, font(24, True), CENTER)
    sheet.row_dimensions[1].height = 31.2
    row = 3
    for label, value in (('Kelompok', group_label(record.kelompok)), ('Jadwal Shift', record.shift),
                         ('Hari / Tanggal', _day(record.date)),
                         ('Waktu Pengecekan', f'{record.check_time:%H:%M} WIB')):
        _write(sheet, row, 'B', 'C', label, font(16, True), SHRINK)
        _write(sheet, row, 'D', 'D', f': {value}', font(16, True), SHRINK)
        sheet.row_dimensions[row].height = 23.4
        row += 1
    for section in record_sections(record):
        row = _answer_table(sheet, section, row + 1)
    last_row = _email_web_signatures(sheet, record, row + 1)
    _page_setup(sheet, 'H', last_row, PageMargins(left=0, right=0, top=0.51, bottom=0, header=0, footer=0))


def _choices(sheet, row, column, values, selected):
    """The options of an answer in two rows, each with a box that is checked for the answer given."""
    for offset, value in enumerate(values):
        box = CHECKED if value == selected else UNCHECKED
        # DejaVu Sans has the boxes: with Calibri, LibreOffice falls back to a color emoji font.
        _write(sheet, row + offset, column, column, f'{value} {box}', Font(name='DejaVu Sans', size=10, bold=True),
               CENTER, border=True)


def _answer_table(sheet, section, row):
    """Email or web table: every item takes two rows (username / password, Y / N, J / BJ or B / BB)."""
    email = section['kind'] == ChecklistSection.Kind.EMAIL
    if section['heading']:
        _write(sheet, row, 'A', 'H', section['heading'],
               font(18 if email else 16, True, italic=True, underline='single'), SHRINK)
        sheet.row_dimensions[row].height = 23.4
        row += 1
    if not email or not section['heading']:  # The email table of the template has no title of its own.
        _write(sheet, row, 'A', 'H', section['title'], font(16, True), SHRINK)
        sheet.row_dimensions[row].height = 21
        row += 1

    header = {'alignment': CENTER, 'border': True, 'fill': HEADER_FILL}
    bold = font(12, True, color=WHITE)
    _write(sheet, row, 'A', 'A', 'No.', bold, last_row=row + 1, **header)
    _write(sheet, row, 'B', 'C', 'Alamat Email' if email else 'Nama Web', bold, last_row=row + 1, **header)
    _write(sheet, row, 'D', 'D', 'Username & Password' if email else 'Alamat Web', bold, last_row=row + 1, **header)
    _write(sheet, row, 'E', 'E', 'Akses Email' if email else 'Akses Web', font(8, True, color=WHITE),
           last_row=row + 1, **header)
    if email:
        _write(sheet, row, 'F', 'F', 'Keterangan\n(judul email jika bisa diakses/tindakan jika tidak bisa diakses)',
               font(11, True, color=WHITE), last_row=row + 1, **header)
    else:
        _write(sheet, row, 'F', 'F', 'Keterangan Informasi', bold, last_row=row + 1, **header)
    _write(sheet, row, 'G', 'H', 'Konfirmasi', bold, **header)
    _write(sheet, row + 1, 'G', 'G', 'J/BJ' if email else 'B/BB', bold, **header)
    _write(sheet, row + 1, 'H', 'H', 'Forward Ke' if email else 'Disampaikan Ke', bold, **header)
    sheet.row_dimensions[row].height = sheet.row_dimensions[row + 1].height = 22 if email else 20.1
    row += 2

    first_item_row = row
    confirm = [value for value, _ in section['confirm_choices']]
    for item in section['rows']:
        top, bottom = row, row + 1
        _write(sheet, top, 'A', 'A', item['no'], font(12), CENTER, border=True, last_row=bottom)
        texts = [(item['info'], _width(sheet, 'F', 'F'), 11), (item['forward_to'], _width(sheet, 'H', 'H'), 11)]
        if email:
            _write(sheet, top, 'B', 'B', item['name'], font(12), CENTER, border=True, last_row=bottom)
            _write(sheet, top, 'D', 'D', f"username : {item['address']}", font(12), LEFT)
            _write(sheet, bottom, 'D', 'D', f"password : {item['password']}", font(12), LEFT)
            _outline(sheet, top, bottom, 'D', 'D')
            texts.append((item['name'], _width(sheet, 'B', 'B'), 12))
        else:
            _write(sheet, top, 'B', 'C', item['name'], font(12), LEFT, border=True, last_row=bottom)
            _write(sheet, top, 'D', 'D', item['address'], font(11), LEFT, border=True, last_row=bottom)
            texts += [(item['name'], _width(sheet, 'B', 'C'), 12), (item['address'], _width(sheet, 'D', 'D'), 11)]
        _choices(sheet, top, 'E', ('Y', 'N'), item['access'])
        _write(sheet, top, 'F', 'F', item['info'] or None, font(11), LEFT, border=True, last_row=bottom)
        _choices(sheet, top, 'G', confirm, item['confirmed'])
        _write(sheet, top, 'H', 'H', item['forward_to'] or None, font(11), LEFT, border=True, last_row=bottom)
        _fit_height(sheet, [top, bottom], texts, 24.9 if not email else 23.4)
        row += 2
    if email and section['rows']:
        _write(sheet, first_item_row, 'C', 'C', section['side_note'] or None, font(10), CENTER, border=True,
               last_row=row - 1)
    return _notes(sheet, section, row, ('B', 'D'), ('E', 'H'))


def _email_web_signatures(sheet, record, row):
    """The signature box at the bottom right: 'Analis on Duty' with the operator's name, and 'Pejabat on Duty'."""
    bold = font(12, True)
    _write(sheet, row + 1, 'G', 'H', f'Jakarta, {format_date_indonesian(record.date)}', bold, CENTER)
    _write(sheet, row + 2, 'F', 'F', 'Analis on Duty', bold, CENTER)
    _write(sheet, row + 2, 'G', 'H', 'Mengetahui,', bold, CENTER)
    _write(sheet, row + 3, 'G', 'H', 'Pejabat on Duty', bold, CENTER)
    _write(sheet, row + 7, 'F', 'F', record.operator.name, font(12, True, underline='single'), CENTER)
    _write(sheet, row + 7, 'G', 'H', '(................................................)', font(12), CENTER)
    for box_row in range(row, row + 9):
        sheet.row_dimensions[box_row].height = 18
    _outline(sheet, row, row + 8, 'F', 'H', style='medium')
    return row + 8


# Checklist SeisComP (Backup), TOAST, Diseminasi, TSP dan WRS NG

def _device_sheet(sheet, record):
    _columns(sheet, DEVICE_WIDTHS)
    _write(sheet, 1, 'A', 'D', record.title, font(16, True), CENTER)
    sheet.row_dimensions[1].height = 22
    bold = font(12, True)
    # The label padding of the template, which lines up the colons.
    _write(sheet, 3, 'A', 'B', f'Kelompok           : {group_label(record.kelompok)}', bold, SHRINK)
    _write(sheet, 4, 'A', 'B', f'Jadwal Shift        : {record.shift}', bold, SHRINK)
    _write(sheet, 5, 'A', 'B', f'Hari/Tanggal       : {_day(record.date)}', bold, SHRINK)
    _write(sheet, 6, 'A', 'B', f'Jam                       : {record.check_time:%H:%M} WIB', bold, SHRINK)
    row = 8
    for section in record_sections(record):
        row = _check_table(sheet, section, row) + 2
    last_row = _device_signatures(sheet, record, row - 1)
    _page_setup(sheet, 'D', last_row, PageMargins(left=0.7, right=0.7, top=0.75, bottom=0.75))


def _check_table(sheet, section, row):
    """Ya/Tidak table: No, what is checked (with its address and password), a check mark under Ya or Tidak."""
    if section['heading']:
        _write(sheet, row, 'A', 'D', section['heading'], font(14, True, italic=True, underline='single'))
        row += 1
    _write(sheet, row, 'A', 'D', section['title'], font(16, True), CENTER, border=True)
    sheet.row_dimensions[row].height = 19.7
    header = font(14, True)
    _write(sheet, row + 1, 'A', 'A', 'No.', header, CENTER, border=True, last_row=row + 2)
    _write(sheet, row + 1, 'B', 'B', 'MONITORING', header, CENTER, border=True, last_row=row + 2)
    _write(sheet, row + 1, 'C', 'D', 'Check List', header, CENTER, border=True)
    _write(sheet, row + 2, 'C', 'C', 'Ya', header, CENTER, border=True)
    _write(sheet, row + 2, 'D', 'D', 'Tidak', header, CENTER, border=True)
    sheet.row_dimensions[row + 1].height = sheet.row_dimensions[row + 2].height = 17.35
    row += 3
    for item in section['rows']:
        text = item['name']
        if item['address']:
            text += f" ({item['address']})"
        if item['password']:
            text += f"\nPassword: {item['password']}"
        _write(sheet, row, 'A', 'A', item['no'], font(11), CENTER, border=True)
        _write(sheet, row, 'B', 'B', text, font(11), LEFT, border=True)
        _write(sheet, row, 'C', 'C', CHECK if item['ok'] is True else None, font(11), CENTER, border=True)
        _write(sheet, row, 'D', 'D', CHECK if item['ok'] is False else None, font(11), CENTER, border=True)
        _fit_height(sheet, [row], [(text, _width(sheet, 'B', 'B'), 11)], 15)
        row += 1
    return _notes(sheet, section, row, ('A', 'B'), ('C', 'D'))


def _device_signatures(sheet, record, row):
    """'Analis On Duty' with the operator's name, and 'Mengetahui, Pejabat On Duty' signed by hand."""
    bold = font(13, True)
    _write(sheet, row, 'C', 'D', 'Mengetahui', bold, CENTER)
    _write(sheet, row + 1, 'B', 'B', 'Analis On Duty,', bold, CENTER)
    _write(sheet, row + 1, 'C', 'D', 'Pejabat On Duty', bold, CENTER)
    _write(sheet, row + 5, 'B', 'B', record.operator.name, font(12, True, underline='single'), CENTER)
    _write(sheet, row + 5, 'C', 'D', '(------------------------------)', font(11, True), CENTER)
    return row + 5
