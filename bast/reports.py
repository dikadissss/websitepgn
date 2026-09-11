import json
import re
from io import StringIO
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

from core.dates import convert_to_indonesian, convert_to_roman, date_from_code, format_date_indonesian, get_hari_indonesia
from core.reports import CENTER, LEFT, THIN_BORDER

TEMPLATE_PATH = Path(__file__).resolve().parent / 'static' / 'bast' / 'BAST.xlsx'
MAX_MEMBERS = 10  # Rows K9:L18 of the template.
EVENTS_FIRST_ROW = 29
DEFAULT_ROW_HEIGHT = 15.75
MMI_CHARS_PER_LINE = 23
MEMBER_PRESENT = re.compile(r'^\s*(hadir\s*$|diganti\b)', re.IGNORECASE)
LEFT_EDGE = Border(left=Side(style='medium'))
RIGHT_EDGE = Border(right=Side(style='medium'))
WRAPPED = Alignment(wrap_text=True, vertical='center')


def parse_members(text):
    """Members are stored as JSON [{"nama", "keterangan"}]; very old records hold one name per line."""
    try:
        return json.loads(text) if text else []
    except json.JSONDecodeError:
        return [{'nama': line.strip(), 'keterangan': ''} for line in text.splitlines() if line.strip()]


def group_label(number):
    return f'{convert_to_roman(number)} ({convert_to_indonesian(number)})'


def export_documents(record):
    """Documents printed for the record, in order (see core.duty.JOBS)."""
    return [build_workbook(record)]


def build_workbook(record):
    workbook = openpyxl.load_workbook(TEMPLATE_PATH)
    sheet = workbook.active
    sheet.title = 'BAST'

    record_date = date_from_code(record.bast_id)
    tanggal = format_date_indonesian(record_date)
    sheet['J4'] = group_label(record.kelompok)
    sheet['J6'] = group_label(record.kel_berikut)
    sheet['N4'] = f': {tanggal}'
    sheet['N5'] = f': {get_hari_indonesia(record_date)}'

    present = 0
    for index, member in enumerate(parse_members(record.member)[:MAX_MEMBERS]):
        note = member.get('keterangan') or ''
        sheet[f'J{9 + index}'] = index + 1
        sheet[f'K{9 + index}'] = member.get('nama')
        sheet[f'L{9 + index}'] = note
        if MEMBER_PRESENT.match(note):
            present += 1
    sheet['L19'] = f'{present}'

    sheet['N6'] = f': {record.waktu_pelaksanaan}'
    sheet['G22'] = f'{record.event_indonesia}'
    sheet['G23'] = f'{record.event_luar}'
    sheet['G24'] = f'{record.event_indonesia + record.event_luar}'
    sheet['L22'] = f': {record.event_dirasakan} event'
    sheet['L23'] = f': {record.event_dikirim} event'
    sheet['E33'] = f'Pukul: {record.waktu_cs}'
    sheet['E34'] = f'IA (549) : Gaps = {record.count_gaps} ; Spike = {record.count_spikes} ; Blank = {record.count_blanks}'
    sheet['E38'] = f'Rp {record.pulsa_poco:,.0f}'.replace(',', '.')
    sheet['E40'] = f'{record.poco_exp.strftime("%d %b %Y")}'
    sheet['G40'] = f'{record.samsung_exp.strftime("%d %b %Y")}'
    sheet['C47'] = f'Jakarta, {tanggal}'
    sheet['C55'] = f'{record.spv}'
    sheet['C56'] = f'NIP. {record.NIP}'
    sheet['D44'] = f'{record.notes}'

    _insert_events(sheet, record.events)
    return workbook


def _insert_events(sheet, events_csv):
    """Insert one bordered row per event above the rest of the form (columns C..P, MMI text wrapped)."""
    events = pd.read_csv(StringIO(events_csv))
    if events.empty:  # insert_rows(amount=0) would erase every row below EVENTS_FIRST_ROW.
        return
    sheet.insert_rows(EVENTS_FIRST_ROW, amount=len(events))

    for sheet_row, row in enumerate(dataframe_to_rows(events, index=False, header=False), EVENTS_FIRST_ROW):
        for column, value in enumerate(row, 3):
            sheet.cell(row=sheet_row, column=column, value=value).alignment = CENTER
        sheet.cell(row=sheet_row, column=2).border = LEFT_EDGE
        sheet.cell(row=sheet_row, column=17).border = RIGHT_EDGE
        sheet.cell(row=sheet_row, column=11).alignment = LEFT

    for sheet_row in range(EVENTS_FIRST_ROW, EVENTS_FIRST_ROW + len(events)):
        mmi = sheet.cell(row=sheet_row, column=12).value
        if pd.notna(mmi):
            if len(str(mmi)) > MMI_CHARS_PER_LINE:
                sheet.row_dimensions[sheet_row].height = DEFAULT_ROW_HEIGHT * (len(str(mmi)) // MMI_CHARS_PER_LINE + 1)
            sheet.cell(row=sheet_row, column=12).alignment = WRAPPED
        else:
            sheet.row_dimensions[sheet_row].height = DEFAULT_ROW_HEIGHT
        for column in range(3, 17):
            sheet.cell(row=sheet_row, column=column).border = THIN_BORDER
