"""Styles and sheet builders shared by the Excel reports of the record apps."""
from io import StringIO

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.dataframe import dataframe_to_rows

from .dates import date_from_code, format_date_indonesian, get_hari_indonesia

CENTER = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center')
THIN = Side(style='thin')
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PREVIOUS_GROUP_FILL = PatternFill(start_color='FFD3D3D3', end_color='FFD3D3D3', fill_type='solid')
SIGNATURE_FONT = Font(name='Calibri', underline='single', size=18, bold=True)

QC_FIRST_ROW = 8
QC_SIGNATURE_COLUMN = 13


def fill_qc_sheet(sheet, record, previous_csv, current_csv, current_label, table_width, left_aligned_column=None):
    """Fill a QC form: header, one grey row per event of the previous group each followed by this group's QC
    row, then the operator's signature block.

    ``table_width`` is the number of bordered columns (from column B); ``current_label`` marks this group's rows.
    """
    record_date = date_from_code(record.code)
    tanggal = format_date_indonesian(record_date)
    sheet['G2'] = ': ' + tanggal
    sheet['G3'] = ': ' + get_hari_indonesia(record_date)
    sheet['G4'] = f': {record.jam_pelaksanaan.strftime("%H:%M")} - selesai'
    sheet['G5'] = f': Kel. {record.kelompok}'

    previous = pd.read_csv(StringIO(previous_csv))
    previous['prev'] = f'Kel. {record.kel_sebelum}'
    current = pd.read_csv(StringIO(current_csv))
    current[current_label] = current_label
    event_count = len(previous)
    if event_count:  # insert_rows(amount=0) would erase every row below QC_FIRST_ROW.
        sheet.insert_rows(QC_FIRST_ROW, amount=event_count * 2)

    for number, row in enumerate(dataframe_to_rows(previous, index=False, header=False), 1):
        sheet_row = number * 2 + 6
        sheet.cell(row=sheet_row, column=2, value=number).alignment = CENTER
        for column, value in enumerate(row, 3):
            cell = sheet.cell(row=sheet_row, column=column, value=value)
            cell.alignment = CENTER
            cell.fill = PREVIOUS_GROUP_FILL
        sheet.merge_cells(start_row=sheet_row, start_column=2, end_row=sheet_row + 1, end_column=2)

    for number, row in enumerate(dataframe_to_rows(current, index=False, header=False), 1):
        for column, value in enumerate(row, 3):
            sheet.cell(row=number * 2 + 7, column=column, value=value).alignment = CENTER

    for sheet_row in range(QC_FIRST_ROW, QC_FIRST_ROW + event_count * 2):
        if left_aligned_column:
            sheet.cell(row=sheet_row, column=left_aligned_column).alignment = LEFT
        for column in range(2, table_width + 2):
            sheet.cell(row=sheet_row, column=column).border = THIN_BORDER
        sheet.row_dimensions[sheet_row].height = 15

    signature_row = QC_FIRST_ROW + event_count * 2
    sheet.cell(row=signature_row + 2, column=QC_SIGNATURE_COLUMN, value=f'Jakarta, {tanggal}')
    sheet.row_dimensions[signature_row + 2].height = 23.5
    sheet.row_dimensions[signature_row + 3].height = 23.5
    sheet.cell(row=signature_row + 8, column=QC_SIGNATURE_COLUMN, value=record.operator.name).font = SIGNATURE_FONT
    sheet.row_dimensions[signature_row + 8].height = 23.5
    sheet.cell(row=signature_row + 9, column=QC_SIGNATURE_COLUMN, value='NIP. ' + record.operator.NIP)
